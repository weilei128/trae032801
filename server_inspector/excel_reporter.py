"""
Excel报表生成模块
用于将服务器巡检结果汇总为Excel报表
"""

import os
from typing import List
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime
from ssh_collector import ServerMetrics


class ExcelReporter:
    """Excel报表生成器"""
    
    def __init__(self, report_path: str):
        self.report_path = report_path
        self.workbook = None
        
        # 定义样式
        self.header_font = Font(bold=True, color="FFFFFF", size=11)
        self.header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        self.header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        
        self.cell_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        self.cell_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # 告警样式
        self.warning_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
        self.warning_font = Font(color="9C0006")
        self.success_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
        self.success_font = Font(color="006100")
        
    def create_workbook(self):
        """创建新的工作簿"""
        self.workbook = Workbook()
        # 删除默认的sheet
        if 'Sheet' in self.workbook.sheetnames:
            self.workbook.remove(self.workbook['Sheet'])
    
    def load_or_create_workbook(self):
        """加载或创建工作簿"""
        if os.path.exists(self.report_path):
            try:
                self.workbook = load_workbook(self.report_path)
            except:
                self.create_workbook()
        else:
            self.create_workbook()
    
    def get_or_create_sheet(self, sheet_name: str):
        """获取或创建工作表"""
        if sheet_name in self.workbook.sheetnames:
            return self.workbook[sheet_name]
        else:
            return self.workbook.create_sheet(title=sheet_name)
    
    def set_column_widths(self, worksheet, widths: dict):
        """设置列宽"""
        for col, width in widths.items():
            worksheet.column_dimensions[col].width = width
    
    def add_summary_sheet(self, all_metrics: List[ServerMetrics]):
        """添加汇总数据表"""
        ws = self.get_or_create_sheet("巡检汇总")
        
        # 清空现有内容
        ws.delete_rows(1, ws.max_row)
        
        # 标题行
        headers = [
            "巡检时间", "服务器名称", "服务器IP", "连接状态",
            "CPU负载(1分钟)", "CPU负载(5分钟)", "CPU负载(15分钟)", "CPU使用率(%)",
            "内存总量(MB)", "内存使用(MB)", "内存剩余(MB)", "内存使用率(%)",
            "磁盘总量(GB)", "磁盘使用(GB)", "磁盘剩余(GB)", "磁盘使用率(%)",
            "TCP监听端口", "总进程数", "运行中进程数",
            "Docker已安装", "Docker版本", "Docker镜像数", 
            "容器总数", "运行中容器", "已停止容器",
            "目标目录存在", "目标目录大小(MB)", "错误信息"
        ]
        
        # 写入标题
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = self.header_font
            cell.fill = self.header_fill
            cell.alignment = self.header_alignment
            cell.border = self.cell_border
        
        # 写入数据
        for row_idx, metrics in enumerate(all_metrics, 2):
            row_data = [
                metrics.timestamp,
                metrics.server_name,
                metrics.server_ip,
                metrics.connection_status,
                metrics.cpu_load_1min,
                metrics.cpu_load_5min,
                metrics.cpu_load_15min,
                metrics.cpu_usage_percent,
                metrics.memory_total_mb,
                metrics.memory_used_mb,
                metrics.memory_free_mb,
                metrics.memory_usage_percent,
                metrics.disk_total_gb,
                metrics.disk_used_gb,
                metrics.disk_free_gb,
                metrics.disk_usage_percent,
                metrics.tcp_listening_ports,
                metrics.total_processes,
                metrics.running_processes,
                "是" if metrics.docker_installed else "否",
                metrics.docker_version,
                metrics.docker_images_count,
                metrics.docker_containers_total,
                metrics.docker_containers_running,
                metrics.docker_containers_stopped,
                "是" if metrics.target_dir_exists else "否",
                metrics.target_dir_size_mb,
                metrics.error_message
            ]
            
            for col_idx, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.alignment = self.cell_alignment
                cell.border = self.cell_border
                
                # 根据数值设置告警样式
                if col_idx == 4:  # 连接状态列
                    if metrics.connection_status == "Success":
                        cell.fill = self.success_fill
                        cell.font = self.success_font
                    elif metrics.connection_status in ["Failed", "Error"]:
                        cell.fill = self.warning_fill
                        cell.font = self.warning_font
                
                # CPU使用率告警 (>80%)
                elif col_idx == 8 and isinstance(value, (int, float)):
                    if value > 80:
                        cell.fill = self.warning_fill
                        cell.font = self.warning_font
                
                # 内存使用率告警 (>85%)
                elif col_idx == 12 and isinstance(value, (int, float)):
                    if value > 85:
                        cell.fill = self.warning_fill
                        cell.font = self.warning_font
                
                # 磁盘使用率告警 (>90%)
                elif col_idx == 16 and isinstance(value, (int, float)):
                    if value > 90:
                        cell.fill = self.warning_fill
                        cell.font = self.warning_font
        
        # 设置列宽
        column_widths = {
            'A': 20, 'B': 18, 'C': 16, 'D': 12,
            'E': 15, 'F': 15, 'G': 16, 'H': 14,
            'I': 15, 'J': 15, 'K': 15, 'L': 15,
            'M': 15, 'N': 15, 'O': 15, 'P': 15,
            'Q': 30, 'R': 12, 'S': 14,
            'T': 14, 'U': 25, 'V': 14,
            'W': 12, 'X': 12, 'Y': 12,
            'Z': 16, 'AA': 18, 'AB': 30
        }
        self.set_column_widths(ws, column_widths)
        
        # 冻结首行
        ws.freeze_panes = 'A2'
    
    def add_statistics_sheet(self, all_metrics: List[ServerMetrics]):
        """添加统计信息表"""
        ws = self.get_or_create_sheet("统计信息")
        ws.delete_rows(1, ws.max_row)
        
        # 统计标题
        ws.cell(row=1, column=1, value="服务器巡检统计信息").font = Font(bold=True, size=14)
        ws.merge_cells('A1:D1')
        
        # 生成时间
        ws.cell(row=2, column=1, value=f"生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        ws.merge_cells('A2:D2')
        
        # 空行
        current_row = 4
        
        # 按服务器分组统计
        servers = {}
        for metric in all_metrics:
            if metric.server_ip not in servers:
                servers[metric.server_ip] = []
            servers[metric.server_ip].append(metric)
        
        for server_ip, metrics_list in servers.items():
            # 服务器标题
            ws.cell(row=current_row, column=1, value=f"服务器: {metrics_list[0].server_name} ({server_ip})").font = Font(bold=True, size=12)
            ws.merge_cells(f'A{current_row}:D{current_row}')
            current_row += 1
            
            # 统计表头
            stat_headers = ["指标", "平均值", "最大值", "最小值"]
            for col, header in enumerate(stat_headers, 1):
                cell = ws.cell(row=current_row, column=col, value=header)
                cell.font = self.header_font
                cell.fill = self.header_fill
                cell.alignment = self.header_alignment
                cell.border = self.cell_border
            current_row += 1
            
            # 计算统计数据
            stats_data = [
                ("CPU负载(1分钟)", [m.cpu_load_1min for m in metrics_list]),
                ("CPU负载(5分钟)", [m.cpu_load_5min for m in metrics_list]),
                ("CPU负载(15分钟)", [m.cpu_load_15min for m in metrics_list]),
                ("CPU使用率(%)", [m.cpu_usage_percent for m in metrics_list]),
                ("内存使用率(%)", [m.memory_usage_percent for m in metrics_list]),
                ("磁盘使用率(%)", [m.disk_usage_percent for m in metrics_list]),
                ("运行中进程数", [m.running_processes for m in metrics_list]),
            ]
            
            for stat_name, values in stats_data:
                if values and any(values):
                    avg_val = sum(values) / len(values)
                    max_val = max(values)
                    min_val = min(values)
                else:
                    avg_val = max_val = min_val = 0
                
                row_data = [stat_name, round(avg_val, 2), max_val, min_val]
                for col, value in enumerate(row_data, 1):
                    cell = ws.cell(row=current_row, column=col, value=value)
                    cell.alignment = self.cell_alignment
                    cell.border = self.cell_border
                current_row += 1
            
            current_row += 2  # 空行
        
        # 设置列宽
        ws.column_dimensions['A'].width = 20
        ws.column_dimensions['B'].width = 15
        ws.column_dimensions['C'].width = 15
        ws.column_dimensions['D'].width = 15
    
    def generate_report(self, all_metrics: List[ServerMetrics]):
        """生成完整的Excel报表"""
        self.load_or_create_workbook()
        
        # 添加汇总表
        self.add_summary_sheet(all_metrics)
        
        # 添加统计表
        self.add_statistics_sheet(all_metrics)
        
        # 保存文件
        os.makedirs(os.path.dirname(self.report_path), exist_ok=True)
        self.workbook.save(self.report_path)
        print(f"报表已生成: {self.report_path}")
        
        return self.report_path
